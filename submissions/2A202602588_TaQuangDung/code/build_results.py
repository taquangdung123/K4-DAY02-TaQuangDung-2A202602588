"""Build the seven rubric workbook sheets from recorded experiments."""
from __future__ import annotations

import json
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[3]
SUB = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import eval as eval_lib

RUNS = SUB / "outputs" / "runs"
PRED = SUB / "predictions"
CLASSES = ["Chinee Apple", "Lantana", "Parkinsonia", "Parthenium", "Prickly Acacia",
           "Rubber Vine", "Siam Weed", "Snake Weed", "Negative"]


def summary(exp_id, seed=0):
    path = RUNS / exp_id / f"seed{seed}" / "summary.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


def metrics(path):
    frame = pd.read_csv(path)
    y = frame.y_true.to_numpy(dtype=int)
    probs = frame[[f"p{i}" for i in range(9)]].to_numpy(dtype=float)
    return eval_lib.compute_metrics(y, probs.argmax(1), probs)


def training_sheet():
    base = summary("B03")
    rows = []
    specs = [("T00", "baseline", "none", "B03")]
    specs += [("T01", "A: initialization", "frozen backbone", "T01"),
              ("T02", "B: augmentation", "color jitter", "T02"),
              ("T03", "C: loss", "label smoothing 0.1", "T03"),
              ("T04", "combination", "color jitter + label smoothing", "T04")]
    for exp_id, axis, change, run_id in specs:
        result = summary(run_id)
        if not result:
            continue
        pred = SUB / "outputs" / "predictions" / f"{run_id}_seed0_val.csv"
        m = metrics(pred)
        rows.append({"exp_id": exp_id, "backbone": "convnext_tiny", "axis": axis,
                     "change_from_T00": change, "seed": 0, "val_macro_f1": result["val_macro_f1"],
                     "val_top1": result["val_top1"],
                     "delta_macro_f1_vs_T00": result["val_macro_f1"] - base["val_macro_f1"],
                     "Chinee_Apple_val_f1": float(m["f1"][0]),
                     "Snake_Weed_val_f1": float(m["f1"][7]),
                     "best_epoch": result["best_epoch"], "epoch_seconds": result["epoch_seconds_mean"],
                     "note": "one seed; val only"})
    return pd.DataFrame(rows)


def inference_sheet(selected):
    folder = SUB / "outputs" / "inference" / selected
    path = folder / "val_methods.csv"
    if not path.exists():
        return pd.DataFrame()
    frame = pd.read_csv(path)
    latency_path = SUB / "outputs" / "inference" / "latency.csv"
    if latency_path.exists():
        latency = pd.read_csv(latency_path).set_index("configuration")
        mapping = {"I00": "B03_fp32_b1", "I01": "B03_flip_b1", "I02": "B03_multiscale_b1",
                   "I03": "B03_flip_b1", "I04": "B03_256_b1", "I05": "B03_B01_ensemble_b1",
                   "I07": "B03_fp32_b1"}
        for field in ("p50", "p95", "p99", "images_per_s"):
            frame[f"latency_{field}"] = [float(latency.loc[mapping[e], field]) if e in mapping else np.nan
                                               for e in frame.exp_id]
        reference = float(latency.loc["B03_fp32_b1", "p50"])
        frame["relative_cost"] = frame.latency_p50 / reference
    return frame


def final_sheets(selected):
    final_rows, class_rows = [], []
    for exp_id in ("T00", "F01"):
        group = []
        for seed in (0, 1, 2):
            path = PRED / f"{exp_id}_seed{seed}_test.csv"
            if not path.exists():
                continue
            m = metrics(path)
            group.append(m)
            val_path = PRED / f"{exp_id}_seed{seed}_val.csv"
            if not val_path.exists() and exp_id == "T00":
                val_path = SUB / "outputs" / "predictions" / f"B03_seed{seed}_val.csv"
            val_f1 = metrics(val_path)["macro_f1"] if val_path.exists() else np.nan
            recipe = (f"{selected}: ConvNeXt-Tiny, 224 px, 12 epochs, "
                      + ("color jitter + label smoothing 0.1" if selected == "T04" else
                         "color jitter" if selected == "T02" else
                         "label smoothing 0.1" if selected == "T03" else
                         "frozen backbone" if selected == "T01" else "basic augmentation + CE")
                      + ", one view + val-fitted temperature") if exp_id == "F01" else (
                          "B03: ConvNeXt-Tiny, 224 px, 12 epochs, basic augmentation + CE, one view")
            final_rows.append({"exp_id": exp_id, "configuration": recipe,
                               "seed": seed, "val_macro_f1": val_f1, "test_macro_f1": m["macro_f1"],
                               "test_top1": m["top1"], "test_ece": m["ece"]})
        if len(group) == 3:
            val_values = [row["val_macro_f1"] for row in final_rows if row["exp_id"] == exp_id and isinstance(row["seed"], int)]
            final_rows.append({"exp_id": exp_id, "configuration": "mean ± sample std (3 seeds)",
                               "seed": "mean", "val_macro_f1": float(np.mean(val_values)),
                               "val_macro_f1_std": float(np.std(val_values, ddof=1)),
                               **{f"{column}_{stat}": float(fn([m[field] for m in group]))
                                  for field, column in (("macro_f1", "test_macro_f1"), ("top1", "test_top1"), ("ece", "test_ece"))
                                  for stat, fn in (("mean", np.mean), ("std", lambda x: np.std(x, ddof=1)))}})
            for i, name in enumerate(CLASSES):
                class_rows.append({"exp_id": exp_id, "class": name, "test_support": int(group[0]["support"][i]),
                                   **{field: float(np.mean([m[field][i] for m in group]))
                                      for field in ("precision", "recall", "f1")},
                                   **{field + "_std": float(np.std([m[field][i] for m in group], ddof=1))
                                      for field in ("precision", "recall", "f1")}})
    return pd.DataFrame(final_rows), pd.DataFrame(class_rows)


def main():
    old = pd.read_excel(SUB / "results.xlsx", sheet_name="Backbones")
    training = training_sheet()
    decision = SUB / "outputs" / "final_selection.json"
    selected = json.loads(decision.read_text(encoding="utf-8"))["selected_training_run"] if decision.exists() else "B03"
    inferred = inference_sheet(selected)
    final, per_class = final_sheets(selected)
    latency_path = SUB / "outputs" / "inference" / "latency.csv"
    latency = pd.read_csv(latency_path) if latency_path.exists() else pd.DataFrame()
    candidates = []
    for _, row in old.iterrows():
        candidates.append({"exp_id": row.exp_id, "stage": "backbone", "val_macro_f1": row.val_macro_f1,
                           "val_top1": row.val_top1, "latency_p50_ms": row.get("batch1_latency_ms_preliminary")})
    for _, row in training.iterrows():
        candidates.append({"exp_id": row.exp_id, "stage": "training", "val_macro_f1": row.val_macro_f1,
                           "val_top1": row.val_top1, "latency_p50_ms": np.nan})
    for _, row in inferred.iterrows():
        candidates.append({"exp_id": row.exp_id, "stage": "inference", "val_macro_f1": row.val_macro_f1,
                           "val_top1": row.val_top1, "latency_p50_ms": row.get("latency_p50", np.nan)})
    top = pd.DataFrame(candidates).sort_values("val_macro_f1", ascending=False).head(10)
    sheets = {"Backbones": old, "Training": training, "Inference": inferred,
              "Final": final, "PerClass": per_class, "Latency": latency, "Summary": top}
    with pd.ExcelWriter(SUB / "results.xlsx", engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="164E63")
            for cells in sheet.columns:
                width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 46)
                sheet.column_dimensions[cells[0].column_letter].width = width
            metric = "val_macro_f1" if "val_macro_f1" in frame.columns else None
            if name == "Final":
                metric = "test_macro_f1_mean"
            if metric and not frame.empty and frame[metric].notna().any():
                winner = int(frame[metric].astype(float).idxmax()) + 2
                if name == "Final":
                    preferred = frame.index[(frame.exp_id == "F01") & (frame.seed == "mean")]
                    if len(preferred):
                        winner = int(preferred[0]) + 2
                for cell in sheet[winner]:
                    cell.fill = PatternFill("solid", fgColor="D1FAE5")
    print({name: len(frame) for name, frame in sheets.items()})


if __name__ == "__main__":
    main()
